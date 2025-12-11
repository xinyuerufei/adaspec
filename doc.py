import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Create a new Document
doc = Document()

# Set default font to satisfy Chinese requirements (SimSun usually)
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_heading(text, level):
    h = doc.add_heading(text, level=level)
    h.style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for run in h.runs:
        run.font.name = 'Times New Roman'

def add_paragraph(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.name = 'Times New Roman'

def add_table(data, header_list):
    table = doc.add_table(rows=1, cols=len(header_list))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(header_list):
        hdr_cells[i].text = header
        # Make header bold
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = 'Times New Roman'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    for row_data in data:
        row_cells = table.add_row().cells
        for i, item in enumerate(row_data):
            row_cells[i].text = str(item)
            # Set font for cells
            for paragraph in row_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.add_paragraph() # Add space after table

# Title
doc.add_heading('实验结果与分析', 0)

# 1. Spherical Aberration
add_heading('1. 球差 (Spherical Aberration)', level=1)
add_heading('1.1 实验数据与计算', level=2)
add_paragraph('选取 φ15mm 孔径（近轴区域）对应的像距作为理想像点位置 L\'_ref。球差计算公式为 δL\' = l\'_avg - L\'_ref。')

data_spherical = [
    ['φ15 (基准)', '148', '149', '151', '149.3', '0.0'],
    ['φ30', '136', '138', '137', '137.0', '-12.3'],
    ['φ40', '127', '127', '128', '127.3', '-22.0'],
    ['φ50', '117', '115', '118', '116.7', '-32.6']
]
header_spherical = ['光阑孔径 (mm)', '测量值1', '测量值2', '测量值3', '像距均值 (mm)', '轴向球差 (mm)']
add_table(data_spherical, header_spherical)

add_heading('1.2 结果分析', level=2)
add_paragraph('实验现象：', bold=True)
add_paragraph('从数据可以看出，随着光阑孔径的增大（φ30 → φ50），光束汇聚后的像距逐渐减小。计算所得的轴向球差均为负值，且孔径越大，球差的绝对值越大（从 12.3mm 增加到 32.6mm）。')
add_paragraph('理论解释：', bold=True)
add_paragraph('本实验使用的是单正透镜。根据球差理论，对于单正透镜，边缘光线（大孔径处）的折射能力比中心近轴光线更强，导致边缘光线与光轴的交点比近轴光线更靠近透镜。这种边缘光线焦距短于中心光线焦距的现象称为负球差。实验结果与理论规律完全一致。')
add_paragraph('[排版建议：请在此处插入原文档中的 图4-1 球差实验装置图]')

# 2. Chromatic Aberration
add_heading('2. 位置色差 (Longitudinal Chromatic Aberration)', level=1)
add_heading('2.1 实验数据与计算', level=2)
add_paragraph('位置色差定义为 F光（蓝光）像距与 C光（红光）像距之差：ΔL\'_FC = l\'_F - l\'_C。')

data_chromatic = [
    ['φ30', '红光 (C)', '146', '145', '145', '145.3', '-16.6'],
    ['', '蓝光 (F)', '129', '128', '129', '128.7', ''],
    ['φ40', '红光 (C)', '134', '134', '134', '134.0', '-11.0'],
    ['', '蓝光 (F)', '123', '122', '124', '123.0', ''],
    ['φ50', '红光 (C)', '122', '121', '121', '121.3', '-8.3'],
    ['', '蓝光 (F)', '112', '113', '114', '113.0', '']
]
header_chromatic = ['光阑孔径', '色光', '测量值1', '测量值2', '测量值3', '均值 (mm)', '位置色差 (mm)']
add_table(data_chromatic, header_chromatic)

add_heading('2.2 结果分析', level=2)
add_paragraph('实验现象：', bold=True)
add_paragraph('在所有孔径设置下，蓝光（短波）的成像位置均比红光（长波）更靠近透镜，计算出的色差均为负值。')
add_paragraph('理论解释：', bold=True)
add_paragraph('光学材料的折射率随波长变化，波长越短折射率越高（正常色散）。因此，蓝光通过透镜后的偏折角大于红光，导致其焦距更短。这验证了单透镜存在显著的轴向色差。此外，实验数据显示随着孔径增大，测得的色差绝对值呈减小趋势。这主要是因为大孔径下存在显著的球差，红光和蓝光的球差量不同，影响了肉眼对“最清晰像面”的判断，导致测量误差增大。')
add_paragraph('[排版建议：请在此处插入原文档中的 图2-7 色差原理图 和 图4-2 位置色差实验装置图]')

# 3. Astigmatism
add_heading('3. 像散 (Astigmatism)', level=1)
add_heading('3.1 实验数据与计算', level=2)
add_paragraph('像散值计算公式：x\'_ts = l\'_t - l\'_s。')

data_astigmatism = [
    ['30°', '79', '124', '-45'],
    ['15°', '122', '143', '-21']
]
header_astigmatism = ['透镜偏转角', '子午像距 l\'_t (mm)', '弧矢像距 l\'_s (mm)', '像散值 x\'_ts (mm)']
add_table(data_astigmatism, header_astigmatism)

add_heading('3.2 结果分析', level=2)
add_paragraph('实验现象：', bold=True)
add_paragraph('当透镜发生偏转（模拟轴外视场）时，光斑无法汇聚成一点，而是先后形成相互垂直的线条（子午焦线和弧矢焦线）。偏转角度从 15° 增至 30° 时，两焦线间的距离（像散值）从 21mm 增大至 45mm。')
add_paragraph('理论解释：', bold=True)
add_paragraph('像散是轴外像差，其大小与视场角（此处由透镜偏转角模拟）有关。随着偏转角增大，入射光束在子午方向和弧矢方向的不对称性增强，导致两个方向的聚焦能力差异变大，从而使像散显著增加。')
add_paragraph('[排版建议：请务必在此处插入 图4-3 像散实验装置图]')

# 4. Field Curvature
add_heading('4. 场曲 (Field Curvature)', level=1)
add_heading('4.1 实验结果', level=2)
add_paragraph('中心清晰像距：184 mm')
add_paragraph('边缘清晰像距：136 mm')
add_paragraph('场曲值 x\'：48 mm')
add_heading('4.2 结果分析', level=2)
add_paragraph('实验中使用十字缝作为物体。调节像屏时发现，十字缝的中心和边缘无法同时清晰成像。当中心成像清晰时，边缘模糊；当边缘成像清晰时，像屏需向透镜方向移动 (184mm → 136mm)。这表明该光学系统的像面不是一个平面，而是一个向透镜方向弯曲的曲面，即存在场曲。')

# 5. Qualitative
add_heading('5. 定性像差观察 (定性实验)', level=1)
add_heading('5.1 彗差 (Coma)', level=2)
data_coma = [
    ['位置1 (光阑在透镜前)', '像斑呈彗星状，尖向外，尾向内', '[请在此处插入光阑在前的实拍彗差图]'],
    ['位置2 (光阑在透镜后)', '像斑呈彗星状，尖向内，尾向外', '[请在此处插入光阑在后的实拍彗差图]']
]
header_coma = ['光阑位置', '现象描述', '图示位置']
add_table(data_coma, header_coma)
add_paragraph('分析：彗差体现了轴外物点宽光束成像的不对称性。光阑位置的改变反转了主光线与边缘光线的相对位置关系，从而导致彗差方向（拖尾方向）的反转。')

add_heading('5.2 倍率色差 (Lateral Color)', level=2)
add_paragraph('光阑在透镜前：观察到环状像边缘呈现“外边红，内边蓝”。')
add_paragraph('光阑在透镜后：观察到环状像边缘呈现“外边蓝，内边红”。')
add_paragraph('[排版建议：建议在此处并列放置两张彩色实拍图（截取圆环边缘部分），并标注位置]')

add_heading('5.3 畸变 (Distortion)', level=2)
data_distortion = [
    ['桶形畸变', '光阑位于透镜与物体之间', '方格网边缘向内收缩，状如桶形', '[插入桶形畸变实拍图]'],
    ['枕形畸变', '光阑位于透镜后方', '方格网边缘向外拉伸，状如枕形', '[插入枕形畸变实拍图]']
]
header_distortion = ['畸变类型', '实验条件', '现象特征', '实验现象图']
add_table(data_distortion, header_distortion)
add_paragraph('分析：畸变是由于主光线的球差导致的，它只改变像的形状而不影响清晰度。光阑位置不同，改变了不同视场主光线在透镜上的入射高度，从而改变了垂轴放大率随视场的变化规律，导致了畸变正负的改变。')

# Save file
file_path = './像差实验结果整理.docx'
doc.save(file_path)
